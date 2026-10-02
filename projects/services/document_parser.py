from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class DocumentExtractionError(ValueError):
    """Raised when a supported report cannot be converted into useful text."""


MAX_EXTRACTED_CHARS = 200_000


def extract_report_text(upload):
    suffix = Path(upload.name).suffix.lower()
    try:
        upload.seek(0)
        raw = upload.read()
        if not raw:
            raise DocumentExtractionError("The selected report is empty.")
        if suffix == ".txt":
            text = raw.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(raw), strict=False)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            document = Document(BytesIO(raw))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
        else:
            raise DocumentExtractionError("Upload a PDF, TXT, or DOCX report.")
    except DocumentExtractionError:
        raise
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError("The text report must use UTF-8 encoding.") from exc
    except Exception as exc:
        raise DocumentExtractionError(
            "We couldn't read this report. Try exporting it again as a PDF, TXT, or DOCX."
        ) from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise DocumentExtractionError(
            "No readable text was found. Upload a text-based report rather than a scanned image."
        )
    return text[:MAX_EXTRACTED_CHARS]